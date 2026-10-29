import csv, glob, re, os, statistics as st
from collections import defaultdict
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Patch
from matplotlib.ticker import PercentFormatter
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT, OUT = r"results-chpc\results", "report"
FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)
OPPS = ["RandomBiasedAI","WorkerRush","LightRush","HeavyRush","RangedRush","EconomyLightRush","mayari","Rojo","Droplet","Coac"]
SCRIPTED = ["RandomBiasedAI","HeavyRush","RangedRush","LightRush","EconomyLightRush","WorkerRush"]
COMP = ["Rojo","mayari","Droplet","Coac"]
NAVY, MID, LIGHT = "#1f3864", "#5b8bd0", "#7fa3d9"

def read(p):
    with open(p, newline="") as f: return list(csv.DictReader(f))
def per_opp(rows, key):
    t = defaultdict(lambda: defaultdict(lambda: [0, 0]))
    for r in rows:
        if (r.get("error") or "").strip(): continue
        c = t[key(r)][r["opponent"]]; c[0] += r["result"].strip().upper().startswith("W"); c[1] += 1
    return {k: {o: 100 * w / g for o, (w, g) in v.items()} for k, v in t.items()}
macro = lambda d: st.mean(d[o] for o in OPPS)

gp = per_opp(read(ROOT + r"\teameval\all-team.csv"), lambda r: r["seed"])
sb = per_opp(read(ROOT + r"\teameval\all-team-sbgp.csv"), lambda r: r["seed"])
hyb = {}
for d in glob.glob(ROOT + r"\adapter-team\sbgp-seed*"):
    files = sorted(glob.glob(d + r"\eval-adapter-final-0300-team-*.csv"))
    if files: hyb[re.search(r"seed(\d+)", d).group(1)] = per_opp(read(files[-1]), lambda r: "x")["x"]
seeds = sorted(s for s in hyb if s in sb)
gpm, sbm = [macro(v) for v in gp.values()], [macro(v) for v in sb.values()]
bm, hm = [macro(sb[s]) for s in seeds], [macro(hyb[s]) for s in seeds]
diffs = [h - b for b, h in zip(bm, hm)]

checks = [("GP runs", len(gpm), 36), ("SBGP runs", len(sbm), 36), ("paired runs", len(seeds), 20),
          ("GP mean", round(st.mean(gpm), 1), 57.5), ("SBGP mean", round(st.mean(sbm), 1), 57.1),
          ("GP sd", round(st.stdev(gpm), 1), 5.9), ("SBGP sd", round(st.stdev(sbm), 1), 3.5),
          ("base mean", round(st.mean(bm), 1), 56.6), ("hybrid mean", round(st.mean(hm), 1), 60.3)]
for name, got, want in checks: print(f"  {name:<12} {got} (frozen {want}) {'OK' if got == want else 'MISMATCH'}")
if any(g != w for _, g, w in checks): raise SystemExit("MISMATCH with frozen numbers: document NOT written.")

# Figure X.1: beeswarm with +-1 SD band
def swarm(ax, vals, y, col):
    placed = []
    for v in sorted(vals):
        k = 0
        while True:
            off = 0 if k == 0 else ((k + 1) // 2) * 0.09 * (1 if k % 2 else -1)
            if all(abs(v - pv) > 0.55 or abs(off - po) > 0.08 for pv, po in placed): break
            k += 1
        placed.append((v, off))
    ax.scatter([v for v, _ in placed], [y + o for _, o in placed], s=40, color=col, zorder=3)
fig, ax = plt.subplots(figsize=(7.5, 2.9))
for y, vals, col in [(1, gpm, NAVY), (0, sbm, MID)]:
    m, s = st.mean(vals), st.stdev(vals)
    ax.add_patch(Rectangle((m - s, y - 0.32), 2 * s, 0.64, color=col, alpha=0.13, lw=0))
    swarm(ax, vals, y, col)
    ax.plot([m, m], [y - 0.38, y + 0.38], color="black", lw=2, zorder=4)
    ax.text(m, y + 0.41, "mean", ha="center", va="bottom", fontsize=8)
ax.set_yticks([1, 0], [f"Ordinary GP\n{st.mean(gpm):.1f}% ± {st.stdev(gpm):.1f}", f"SBGP\n{st.mean(sbm):.1f}% ± {st.stdev(sbm):.1f}"])
ax.set_ylim(-0.6, 1.7); ax.set_xlim(45, 70); ax.xaxis.set_major_formatter(PercentFormatter(decimals=0))
ax.grid(axis="x", alpha=0.4); ax.tick_params(axis="y", length=0)
for sp in ax.spines.values(): sp.set_visible(False)
ax.set_xlabel("Macro win rate per run (10 opponents, 9 unseen maps); shaded band = ±1 SD", fontsize=9)
fig.tight_layout(); f1 = os.path.join(FIG, "fig_x1_gp_vs_sbgp.png"); fig.savefig(f1, dpi=300); plt.close(fig)

# Figure X.2: paired base vs hybrid
fig, ax = plt.subplots(figsize=(4.8, 4.2))
for b, h in zip(bm, hm):
    col = "tab:green" if h > b + 0.05 else ("tab:red" if h < b - 0.05 else "grey")
    ax.plot([0, 1], [b, h], marker="o", ms=4, color=col, alpha=0.7)
ax.plot([0, 1], [st.mean(bm), st.mean(hm)], marker="D", ms=7, color=NAVY, lw=2.5, label="Mean")
ax.set_xticks([0, 1], ["SBGP bot", "SBGP bot + adapter"]); ax.set_xlim(-0.3, 1.3)
ax.yaxis.set_major_formatter(PercentFormatter(decimals=0)); ax.set_ylabel("Macro win rate")
ax.legend(loc="lower right", frameon=False); ax.grid(axis="y", alpha=0.4)
for sp in ("top", "right"): ax.spines[sp].set_visible(False)
fig.tight_layout(); f2 = os.path.join(FIG, "fig_x2_adapter_paired.png"); fig.savefig(f2, dpi=300); plt.close(fig)

# Figure X.3: per-opponent paired bars, poster style
def mean_opp(src, o): return st.mean(src[s][o] for s in seeds)
rows = [("bar", "Overall")] + [("hdr", "Scripted bots")] + [("bar", o) for o in SCRIPTED] + [("hdr", "Competition bots")] + [("bar", o) for o in COMP]
ys, y = [], 0.0
for kind, _ in rows: ys.append(y); y -= 1.0 if kind == "bar" else 0.65
fig, ax = plt.subplots(figsize=(7.2, 0.5 * len(rows) + 1.0))
for (kind, name), yy in zip(rows, ys):
    if kind != "bar": continue
    b = st.mean(bm) if name == "Overall" else mean_opp(sb, name)
    h = st.mean(hm) if name == "Overall" else mean_opp(hyb, name)
    fmt = "{:.1f}%" if name == "Overall" else "{:.0f}%"
    ax.barh(yy + 0.19, b, height=0.36, color=LIGHT); ax.barh(yy - 0.19, h, height=0.36, color=NAVY)
    ax.text(b + 0.8, yy + 0.19, fmt.format(b), va="center", fontsize=8)
    ax.text(h + 0.8, yy - 0.19, fmt.format(h), va="center", fontsize=8, fontweight="bold" if name == "Overall" else "normal")
ax.set_yticks(ys, [n for _, n in rows])
for tl, (kind, name) in zip(ax.get_yticklabels(), rows):
    if kind == "hdr" or name == "Overall": tl.set_fontweight("bold"); tl.set_color(NAVY)
ax.set_xlim(0, 100); ax.xaxis.set_major_formatter(PercentFormatter(decimals=0)); ax.grid(axis="x", alpha=0.4)
ax.tick_params(axis="y", length=0)
for sp in ax.spines.values(): sp.set_visible(False)
ax.legend(handles=[Patch(color=LIGHT, label="SBGP bot"), Patch(color=NAVY, label="SBGP bot + adapter")],
          loc="lower center", bbox_to_anchor=(0.5, 1.0), ncol=2, frameon=False)
fig.tight_layout(); f3 = os.path.join(FIG, "fig_x3_adapter_per_opponent.png"); fig.savefig(f3, dpi=300); plt.close(fig)

# Document
doc = Document()
def P(text): return doc.add_paragraph(text)
def PB(lead, text):
    p = doc.add_paragraph(); p.add_run(lead).bold = True; p.add_run(text); return p
def CAP(lead, text):
    p = doc.add_paragraph(); r = p.add_run(lead); r.bold = True; r.font.size = Pt(9)
    r2 = p.add_run(" " + text); r2.font.size = Pt(9); return p
def FIGP(path, width):
    doc.add_picture(path, width=Inches(width)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
def TABLE(header, body):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header): c.text = ""; c.paragraphs[0].add_run(h).bold = True
    for row in body:
        cells = t.add_row().cells
        for c, v in zip(cells, row): c.text = str(v)
    doc.add_paragraph()
def PH(text):
    p = doc.add_paragraph(); r = p.add_run(text); r.italic = True

doc.add_heading("X Results", level=1)
P("This chapter reports the performance of the evolved bots under the shared evaluation protocol described in Section [Experimental Setup]. Each bot played every opponent in the ten-bot panel on each of the nine evaluation maps, from both starting positions, with a limit of 2000 game cycles. The four deterministic opponents (WorkerRush, LightRush, HeavyRush and RangedRush) were played once per map and side, since repeated games against them are identical. The six stochastic opponents were played five times per map and side, giving 612 games per evaluated bot. None of the evaluation maps was used during training or selection. Six of the ten opponents (HeavyRush, RangedRush, mayari, Rojo, Droplet and Coac) were never encountered during evolution or adapter training.")
P("The primary measure is the macro win rate: the win rate against each opponent, averaged with every opponent weighted equally. Draws count as non-wins. Each evolutionary run contributes one value, the macro win rate of its best-of-run individual, evaluated in its original, unpruned form.")

doc.add_heading("X.1 Ordinary GP and structure-based GP", level=2)
P("Both methods were run 36 times with independent random seeds and identical settings apart from the breeding procedure: the same function and terminal sets, fitness function, training maps, training opponents, parsimony pressure and generation budget. Table X.1 summarises the distribution of macro win rates across runs, and Figure X.1 shows every run.")
CAP("Table X.1.", "Macro win rate across 36 independent runs per method (%).")
f = lambda v: f"{v:.1f}"
TABLE(["", "Ordinary GP", "SBGP"], [
    ["Mean", f(st.mean(gpm)), f(st.mean(sbm))], ["Standard deviation", f(st.stdev(gpm)), f(st.stdev(sbm))],
    ["Median", f(st.median(gpm)), f(st.median(sbm))], ["Minimum", f(min(gpm)), f(min(sbm))], ["Maximum", f(max(gpm)), f(max(sbm))]])
FIGP(f1, 6.3)
CAP("Figure X.1.", "Macro win rate of each of the 36 runs per method. Each dot is one run; the black line marks the mean and the shaded band ±1 standard deviation.")
P("The two methods did not differ in central performance (Mann-Whitney U = 641, p ≈ 0.94). They did differ in reliability: SBGP's results were significantly less variable across runs (Brown-Forsythe F(1, 70) = 14.42, p < 0.001). The variance of ordinary GP was 2.8 times that of SBGP, and SBGP's weakest run (52.0%) was 3.6 percentage points better than ordinary GP's (48.4%). Both methods' best runs reached about 67%. Structure-based search therefore did not raise the typical or the best result, but it made a poor result less likely.")
P("Table X.2 breaks the mean win rate down by opponent.")
CAP("Table X.2.", "Mean win rate per opponent across 36 runs (%).")
TABLE(["Opponent", "Ordinary GP", "SBGP"], [[o, f(st.mean(gp[s][o] for s in gp)), f(st.mean(sb[s][o] for s in sb))] for o in OPPS])
P("The two methods' profiles are closely similar. Both bots won most games against the built-in rush bots that build an army (LightRush, HeavyRush, RangedRush and EconomyLightRush), including HeavyRush and RangedRush, which were never seen during training. Against the competition bots (mayari, Rojo, Droplet and Coac) both methods won between roughly 30% and 45% of games. WorkerRush remained the hardest built-in opponent for both, at about one win in three. Per-opponent differences between the methods were not tested individually, and none exceeds 3.5 percentage points.")
PB("Decision time. ", "In 68 of the 72 runs, the slowest single decision during evaluation took at most 90 ms, within the 100 ms classic-track budget. Four runs exceeded it in their worst case: two ordinary GP runs (192 ms and 115 ms) and two SBGP runs (231 ms and 127 ms). These maxima were recorded while 24 evaluations ran in parallel on one cluster node, so they reflect worst-case timing under heavy machine load rather than typical behaviour.")
p = PB("Training cost. ", ""); r = p.add_run("[Training-time distribution per platform (laptop and CHPC cluster) to be inserted. Runs that shared a CPU with other jobs are excluded or marked.]"); r.italic = True

doc.add_heading("X.2 Structure-based GP with the adapter", level=2)
P("To measure the effect of the adapter, the 20 SBGP runs performed on the cluster were each wrapped in an adapter trained for a fixed budget of 300 PPO iterations, using the one-hot base distribution and τ = 0.2. The budget and settings were fixed before training. The final checkpoint of every run was evaluated, with no checkpoint selection, and all 20 runs are reported. During evaluation the hybrid played the highest-probability action (argmax). Because every hybrid has its own base bot, the comparison is paired: each hybrid is compared with the run it was built on.")
up = sum(d > 0.05 for d in diffs); eq = sum(abs(d) <= 0.05 for d in diffs); dn = sum(d < -0.05 for d in diffs)
CAP("Table X.3.", "Paired comparison of 20 SBGP base bots and their adapter-wrapped versions.")
TABLE(["", "Base", "Hybrid"], [["Mean macro win rate", f"{st.mean(bm):.1f}%", f"{st.mean(hm):.1f}%"],
    ["Mean difference", "", f"{st.mean(diffs):+.1f} points"], ["Median difference", "", f"{st.median(diffs):+.1f} points"],
    ["Runs improved / unchanged / worse", "", f"{up} / {eq} / {dn}"]])
FIGP(f2, 4.3)
CAP("Figure X.2.", "Macro win rate of each SBGP base bot and its adapter-wrapped version. Green lines mark runs that improved, red lines runs that got worse, and grey the unchanged run; the dark line connects the means.")
P("The adapter improved the macro win rate in 17 of 20 runs. The two decreases were small: 2.6 and 0.4 points. A Wilcoxon signed-rank test on the paired differences (the unchanged pair excluded, n = 19) found the improvement significant (W = 9, p < 0.001).")
P("The adapters changed the base bot's behaviour substantially: across the 20 runs, the hybrid's chosen action differed from the base bot's on between 13.5% and 24.0% of decisions. In the one run whose win rate was unchanged (67.3% in both cases), the hybrid still played differently from its base, so the identical win rate is a coincidence rather than evidence of an inactive adapter. As a check on the evaluation software, the base bots were also played through the adapter evaluation pipeline with the adapter disabled: for three runs, this reproduced the original results to within 0.5 percentage points (a mean difference of 0.07 points), attributable to the stochastic opponents.")
P("Table X.4 and Figure X.3 show the effect of the adapter per opponent. Because ten separate tests were made, the p-values were adjusted with the Holm correction.")
CAP("Table X.4.", "Mean win rate per opponent for the 20 paired runs (%), with Wilcoxon signed-rank tests on the paired differences.")
dl = {r["opponent"]: r for r in read(r"analysis\hybrid_deltas.csv")}
pf = lambda p: f"{p:.3f}" if p < 0.1 else f"{p:.2f}"
TABLE(["Opponent", "Base", "Hybrid", "Difference", "Improved / unchanged / worse", "p (Holm)"],
      [[o, f(float(dl[o]["base_mean"])), f(float(dl[o]["hybrid_mean"])), f"{float(dl[o]['delta_mean']):+.1f}",
        f"{dl[o]['n_up']} / {dl[o]['n_equal']} / {dl[o]['n_down']}", pf(float(dl[o]["p_holm"]))] for o in OPPS])
FIGP(f3, 6.0)
CAP("Figure X.3.", "Mean win rate per opponent for the 20 SBGP base bots and their adapter-wrapped versions, on the nine unseen evaluation maps.")
P("The mean difference was positive against all ten opponents, so the adapter improved the bot without trading performance against one opponent for another. After correction, the improvement was significant against five opponents: Droplet (+7.7 points), WorkerRush (+7.5), mayari (+4.7), RangedRush (+3.6) and RandomBiasedAI (+2.9). Two of the largest gains came against the competition bots Droplet and mayari, which the adapter never encountered during training. The smallest change was against HeavyRush, against which the base bots already won about 80% of games and half of the runs were unchanged.")
PB("Decision time. ", "The adapter's additional computation per decision step was measured separately on three of the hybrids, single-threaded on a cluster CPU, over about 22,000 decision steps against WorkerRush on the evaluation maps. The adapter added a mean of 2.6 to 3.0 ms per step, and at most 4.4 ms at the 99th percentile, well within the 100 ms budget, even with the base bot's own decision time added. Isolated spikes above 100 ms occurred in 7 of the 22,237 steps (0.03%, at most 211 ms), together with a one-off start-up cost on the first call of each process. As with the base bots, the evaluation games were not forced to stay within the time budget.")

doc.add_heading("X.3 Comparison across the team's approaches", level=2)
PH("[Side-by-side comparison of all team members' configurations, once Shy's evaluation and the GP + adapter results are available. Where protocols differ, for example in training maps, the differences are stated alongside the numbers.]")
doc.add_heading("X.4 Comparison with published results", level=2)
PH("[About three published microRTS results, compared as reported, with the differences in maps, unit table, cycle limits and time budgets stated.]")

out = os.path.join(OUT, "results-chapter.docx"); doc.save(out); print("wrote", out, "and figures in", FIG)