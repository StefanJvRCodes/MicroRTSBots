"""
make_report_docx.py  -  builds Part H (Heinrich Niebuhr) of the research report.
Version 3 (2026-10-28): includes both rounds of review fixes. No patch files needed.

Run from the repo root:   python analysis\\make_report_docx.py
Needs: matplotlib, python-docx (both installed on the laptop).

What it does:
  1. Runs analysis/make_results_docx.py (re-checks the 9 frozen numbers, regenerates
     the three result figures and results-chapter.docx). Stops if any check fails.
  2. Draws the method diagrams (report/figures/diagrams/*.png).
  3. Writes report/part-H-heinrich.docx: the full Part H.
[[TODO]] and [[VERIFY]] markers are highlighted yellow in the document.
"""
import os, re, runpy, statistics as st
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle, Circle
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

# ---------------------------------------------------------------- 1. data + checks
R = runpy.run_path(os.path.join("analysis", "make_results_docx.py"))
gp, sb, hyb, seeds = R["gp"], R["sb"], R["hyb"], R["seeds"]
gpm, sbm, bm, hm, diffs = R["gpm"], R["sbm"], R["bm"], R["hm"], R["diffs"]
OPPS, dl = R["OPPS"], R["dl"]
F1, F2, F3 = R["f1"], R["f2"], R["f3"]
NAVY = R["NAVY"]
PALE, ACC = "#e8eef8", "#fdf2d0"
SHOW = {"mayari": "Mayari", "Coac": "CoacAI"}
show = lambda o: SHOW.get(o, o)
f1d = lambda v: f"{v:.1f}"

OUT = "report"
DIA = os.path.join(OUT, "figures", "diagrams"); os.makedirs(DIA, exist_ok=True)

# ---------------------------------------------------------------- 2. diagram helpers
def _edge(b, toward):
    x, y, w, h = b
    dx, dy = toward[0] - x, toward[1] - y
    if dx == 0 and dy == 0: return x, y
    tx = (w / 2) / abs(dx) if dx else float("inf")
    ty = (h / 2) / abs(dy) if dy else float("inf")
    t = min(tx, ty)
    return x + dx * t, y + dy * t

def _save(fig, name):
    fig.tight_layout(); p = os.path.join(DIA, name + ".png")
    fig.savefig(p, dpi=300); plt.close(fig); return p

def boxflow(name, boxes, arrows, size=(7.5, 3.2), xlim=(0, 10), ylim=(0, 4), extra=None):
    """boxes: (id, x, y, w, h, text[, colour]); arrows: (from, to[, label[, rad[, (lx, ly)]]])"""
    fig, ax = plt.subplots(figsize=size)
    ax.set_xlim(*xlim); ax.set_ylim(*ylim); ax.axis("off")
    if extra: extra(ax)
    pos = {}
    for b in boxes:
        bid, x, y, w, h, text = b[:6]
        fc = b[6] if len(b) > 6 else PALE
        ax.add_patch(FancyBboxPatch((x - w / 2, y - h / 2), w, h,
                     boxstyle="round,pad=0,rounding_size=0.12", fc=fc, ec=NAVY, lw=1.2, zorder=2))
        ax.text(x, y, text, ha="center", va="center", fontsize=8.5, zorder=3)
        pos[bid] = (x, y, w, h)
    for a in arrows:
        s, t = a[0], a[1]
        lab = a[2] if len(a) > 2 else ""
        rad = a[3] if len(a) > 3 else 0.0
        lp = a[4] if len(a) > 4 else None
        p0 = _edge(pos[s], pos[t][:2]); p1 = _edge(pos[t], pos[s][:2])
        ax.annotate("", xy=p1, xytext=p0, zorder=1,
                    arrowprops=dict(arrowstyle="-|>", color=NAVY, lw=1.3, shrinkA=2, shrinkB=2,
                                    connectionstyle=f"arc3,rad={rad}"))
        if lab:
            mx, my = lp if lp else ((p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2)
            ax.text(mx, my, lab, ha="center", va="center", fontsize=7.5, color="#333333",
                    bbox=dict(fc="white", ec="none", pad=1.5), zorder=4)
    return _save(fig, name)

# ---------------------------------------------------------------- 3. diagrams
def d_configs():
    return boxflow("h01_configurations", [
        ("gp", 2, 3.1, 3.0, 0.8, "Ordinary GP\n(canonical Koza GP)"),
        ("sb", 2, 1.9, 3.0, 0.8, "Structure-based GP\n(SBGP)"),
        ("hy", 2, 0.55, 3.0, 0.8, "SBGP + Adapter-RL\n(PPO-trained)", ACC),
        ("ev", 8, 1.9, 2.8, 1.5, "Evaluation\n9 unseen maps\n× 10 opponents\n× 2 sides"),
    ], [
        ("gp", "ev", "same primitives, fitness,\nmaps and parsimony"),
        ("sb", "ev", "changes only breeding"),
        ("sb", "hy", "frozen best-of-run bot", 0, (3.2, 1.22)),
        ("hy", "ev", "adds a learned correction"),
    ], size=(7.5, 3.3), ylim=(0, 3.7))

def d_decision():
    return boxflow("h02_decision_loop", [
        ("u", 3.5, 5.2, 4.2, 0.6, "Idle unit u (checked every game cycle)"),
        ("leg", 3.5, 4.2, 4.2, 0.6, "Engine lists the legal actions a1 … ak"),
        ("sc", 3.5, 3.2, 4.2, 0.6, "Evolved tree scores each action → s1 … sk"),
        ("pk", 3.5, 2.2, 4.2, 0.6, "Try actions from the highest score down"),
        ("af", 3.5, 1.15, 4.2, 0.8, "First affordable action is chosen\n(nothing affordable → wait)"),
        ("hk", 8.3, 3.2, 2.9, 1.0, "Score vector:\nthe hook the adapter\nlater builds on", ACC),
    ], [("u", "leg"), ("leg", "sc"), ("sc", "pk"), ("pk", "af"), ("sc", "hk")],
        size=(7.2, 4.0), ylim=(0.5, 5.7))

def d_terminals():
    fig, ax = plt.subplots(figsize=(7.4, 2.8)); ax.axis("off"); ax.set_aspect("equal")
    ax.set_xlim(-0.2, 11.6); ax.set_ylim(-0.3, 3.3)
    for i in range(4):
        for j in range(3):
            ax.add_patch(Rectangle((i, j), 1, 1, fc="white", ec="#aaaaaa"))
    ax.text(1.5, 1.5, "W", ha="center", va="center", fontsize=15, color=NAVY, weight="bold")
    ax.text(3.5, 2.5, "E", ha="center", va="center", fontsize=15, color="tab:red", weight="bold")
    for (dx, dy, lab, good) in [(0, 1, "UP", 1), (1, 0, "RIGHT", 1), (0, -1, "DOWN", 0), (-1, 0, "LEFT", 0)]:
        col = "tab:green" if good else "grey"
        ax.annotate("", xy=(1.5 + 0.62 * dx, 1.5 + 0.62 * dy), xytext=(1.5 + 0.22 * dx, 1.5 + 0.22 * dy),
                    arrowprops=dict(arrowstyle="-|>", color=col, lw=1.6))
        ax.text(1.5 + 0.8 * dx, 1.5 + 0.8 * dy, lab, ha="center", va="center", fontsize=6.5, color=col)
    ax.text(4.6, 2.35, "movesTowardEnemy\nUP 1   RIGHT 1   DOWN 0   LEFT 0\n→ differs between moves: can steer the unit",
            fontsize=8.5, va="center")
    ax.text(4.6, 0.65, "distEnemyBase\nUP d   RIGHT d   DOWN d   LEFT d\n→ the same for every move: cannot change the choice",
            fontsize=8.5, va="center")
    return _save(fig, "h03_action_discriminating")

def d_roles():
    return boxflow("h04_role_trees", [
        ("u", 1.3, 2.0, 2.0, 0.7, "Idle unit"),
        ("r", 4.0, 2.0, 1.6, 0.7, "Role?"),
        ("t0", 7.6, 3.45, 3.0, 0.7, "Tree 0: worker"),
        ("t1", 7.6, 2.5, 3.0, 0.7, "Tree 1: base"),
        ("t2", 7.6, 1.55, 3.0, 0.7, "Tree 2: barracks"),
        ("t3", 7.6, 0.6, 3.0, 0.7, "Tree 3: combat (Light / Heavy / Ranged)"),
    ], [("u", "r"), ("r", "t0"), ("r", "t1"), ("r", "t2"), ("r", "t3")],
        size=(7.2, 3.0), ylim=(0.1, 4.0))

def d_bands():
    fig, ax = plt.subplots(figsize=(7.2, 2.0)); ax.axis("off")
    ax.set_xlim(-0.1, 3.1); ax.set_ylim(-1.0, 1.2)
    for i, (lab, col, desc) in enumerate([("WIN", "#cfe3cf", "faster win = lower"),
                                          ("DRAW", "#f5e6b8", "better play = lower"),
                                          ("LOSS", "#f2c9c4", "longer survival and\nbetter play = lower")]):
        ax.add_patch(Rectangle((i, 0), 1, 0.6, fc=col, ec=NAVY))
        ax.text(i + 0.5, 0.3, lab, ha="center", va="center", weight="bold")
        ax.text(i + 0.5, -0.15, desc, ha="center", va="top", fontsize=8)
    for v in range(4): ax.text(v, 0.7, str(v), ha="center", va="bottom", fontsize=9)
    ax.text(1.5, 1.05, "Per-game score (lower is better)", ha="center", fontsize=9, color=NAVY)
    return _save(fig, "h05_fitness_bands")

def d_sbgp():
    return boxflow("h06_sbgp_cycle", [
        ("fg", 5, 3.7, 5.4, 0.8, "First global search\ncanonical GP for Sg generations"),
        ("lo", 2.2, 1.3, 3.8, 1.2, "Local search\noperators act below level Dg\nuntil no improvement\nfor Wg generations"),
        ("gl", 7.8, 1.3, 3.8, 1.4, "Global search\noperators act within the top Dg\nlevels for Sg generations; reject\noffspring with similarity ≥ Tg\nto any earlier area"),
    ], [
        ("fg", "lo", "freeze top Dg levels\nof the best → area #1", 0, (2.6, 2.75)),
        ("lo", "gl", "stagnation", -0.3, (5, 1.95)),
        ("gl", "lo", "best dissimilar\nindividual → new area", -0.3, (5, 0.6)),
    ], size=(7.5, 3.6), ylim=(-0.1, 4.3))

def d_similarity():
    fig, ax = plt.subplots(figsize=(7.2, 2.8)); ax.axis("off"); ax.set_aspect("equal")
    ax.set_xlim(-2.2, 7.8); ax.set_ylim(-0.6, 3.0)
    def tree(ox, title, labels, match):
        pos = {"r": (ox, 2), "m": (ox - 0.8, 1), "n": (ox + 0.8, 1), "a": (ox - 1.3, 0), "b": (ox - 0.3, 0)}
        for p, c in [("r", "m"), ("r", "n"), ("m", "a"), ("m", "b")]:
            ax.plot([pos[p][0], pos[c][0]], [pos[p][1], pos[c][1]], color="#777777", zorder=1)
        for k, (x, y) in pos.items():
            fc = {True: "#cfe3cf", False: "#f2c9c4", None: PALE}[match.get(k)]
            ax.add_patch(Circle((x, y), 0.3, fc=fc, ec=NAVY, zorder=2))
            ax.text(x, y, labels[k], ha="center", va="center", fontsize=10, zorder=3)
        ax.text(ox, 2.6, title, ha="center", fontsize=9, color=NAVY, weight="bold")
    tree(0, "Stored global area", {"r": "+", "m": "*", "n": "x", "a": "a", "b": "b"}, {})
    tree(4.2, "Offspring  (similarity = 3)", {"r": "+", "m": "*", "n": "y", "a": "a", "b": "c"},
         {"r": True, "m": True, "a": True, "b": False, "n": False})
    ax.text(5.9, 0.1, "green: node and all its\nancestors match (counted)\nred: no match", fontsize=7.5, va="center")
    return _save(fig, "h07_similarity_index")

def d_dataflow():
    def regions(ax):
        for (x0, x1, title) in [(0.2, 4.6, "JAVA: game engine + frozen evolved bot"),
                                (5.4, 9.95, "PYTHON: PyTorch adapter")]:
            ax.add_patch(Rectangle((x0, 0.45), x1 - x0, 4.45, fc="#f7f7f7", ec="#bbbbbb", lw=1, zorder=0))
            ax.text((x0 + x1) / 2, 4.72, title, ha="center", fontsize=8.5, weight="bold", color=NAVY)
    return boxflow("h08_adapter_dataflow", [
        ("gs", 2.4, 4.1, 3.8, 0.55, "Game state"),
        ("ob", 2.4, 3.2, 3.8, 0.6, "29 feature planes + legal-action mask"),
        ("bs", 2.4, 2.25, 3.8, 0.7, "Frozen SBGP bot proposes\none action per idle unit"),
        ("en", 2.4, 1.05, 3.8, 0.8, "Engine applies the chosen actions\n(base action if unaffordable)"),
        ("nt", 7.7, 3.2, 4.0, 0.7, "Adapter network\n→ 86 adjustment logits per cell + value"),
        ("bl", 7.7, 2.25, 4.0, 0.55, "Base logits = one-hot / τ"),
        ("cb", 7.7, 1.05, 4.0, 0.8, "p = softmax(base + adj), legal actions only\nsample (training) / argmax (evaluation)", ACC),
    ], [
        ("gs", "ob"), ("ob", "nt", "observation (JPype)"), ("bs", "bl", "base action"),
        ("nt", "cb", "", 0, None), ("bl", "cb"), ("cb", "en", "action codes"),
    ], size=(8.0, 4.0), xlim=(0, 10.1), ylim=(0.3, 5.0), extra=regions)

def d_network():
    return boxflow("h09_adapter_network", [
        ("in", 1.0, 2.0, 1.7, 0.9, "Observation\n29 × 16 × 16"),
        ("c", 3.2, 2.0, 1.9, 0.9, "3 × conv 3×3\n64 channels"),
        ("g", 5.45, 2.0, 1.9, 0.9, "Global context\n(pooled, broadcast)"),
        ("m", 7.4, 2.0, 1.4, 0.9, "Conv mix"),
        ("p", 9.5, 3.15, 2.1, 0.8, "1×1 conv → 86 logits\nper cell (zero-init)", ACC),
        ("v", 9.5, 0.85, 2.1, 0.8, "Mean+max pool\n→ MLP → value"),
    ], [("in", "c"), ("c", "g"), ("g", "m"), ("m", "p", "policy"), ("m", "v", "value")],
        size=(8.0, 3.0), xlim=(0, 10.7), ylim=(0.2, 3.8))

def d_workflow():
    return boxflow("h10_workflow", [
        ("cfg", 1.1, 3.0, 1.9, 0.75, "Config\n(.params)"),
        ("evo", 3.7, 3.0, 2.5, 0.75, "ECJ evolution\n(laptop / CHPC)"),
        ("ind", 6.6, 3.0, 2.4, 0.75, "Best-of-run\n.ind + .params"),
        ("ad", 6.6, 1.4, 2.4, 0.75, "Adapter training\n(PPO, CHPC)"),
        ("te", 9.6, 3.0, 2.6, 0.95, "TeamEval\n9 maps × 10 bots × 2 sides\n612 games per bot"),
        ("ae", 9.6, 1.4, 2.6, 0.75, "Adapter evaluation\n(same 612 games)"),
        ("cs", 12.4, 2.2, 1.9, 0.9, "Per-game CSV\n→ statistics", ACC),
    ], [("cfg", "evo"), ("evo", "ind"), ("ind", "te"), ("ind", "ad"), ("ad", "ae"),
        ("te", "cs"), ("ae", "cs")], size=(9.0, 2.9), xlim=(0, 13.5), ylim=(0.8, 3.7))

DG = {k: fn() for k, fn in [("configs", d_configs), ("decision", d_decision), ("terminals", d_terminals),
                            ("roles", d_roles), ("bands", d_bands), ("sbgp", d_sbgp),
                            ("similarity", d_similarity), ("dataflow", d_dataflow),
                            ("network", d_network), ("workflow", d_workflow)]}

# ---------------------------------------------------------------- 4. document helpers
doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(11)
MARK = re.compile(r"(\[\[.*?\]\])")

def runs(p, text, bold=False, italic=False, size=None):
    for part in MARK.split(text):
        if not part: continue
        r = p.add_run(part); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if part.startswith("[["): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p
def H(text, level): doc.add_heading(text, level=level)
def P(text): return runs(doc.add_paragraph(), text)
def PB(lead, text):
    p = doc.add_paragraph(); runs(p, lead, bold=True); return runs(p, text)
def NOTE(text): return runs(doc.add_paragraph(), text, italic=True)
def BUL(items):
    for it in items: runs(doc.add_paragraph(style="List Bullet"), it)
def CODE(text):
    p = doc.add_paragraph(); r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(9)
def CAP(lead, text):
    p = doc.add_paragraph(); runs(p, lead, bold=True, size=9); runs(p, " " + text, size=9)
def FIG(path, width, lead, text):
    doc.add_picture(path, width=Inches(width)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    CAP(lead, text)
def TABLE(lead, caption, header, body):
    CAP(lead, caption)
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header):
        c.text = ""; runs(c.paragraphs[0], h, bold=True, size=9.5)
    for row in body:
        for c, v in zip(t.add_row().cells, row):
            c.text = ""; runs(c.paragraphs[0], str(v), size=9.5)
    doc.add_paragraph()
def EQ():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for txt, sub in [("p(a", 0), ("i", 1), (") = softmax( base", 0), ("i", 1), (" / τ  +  adj", 0),
                     ("i", 1), ("(s) )", 0)]:
        r = p.add_run(txt); r.font.subscript = bool(sub); r.italic = True

# ---------------------------------------------------------------- 5. content
box = doc.add_table(rows=1, cols=1); box.style = "Table Grid"
c = box.rows[0].cells[0]; c.text = ""
runs(c.paragraphs[0], "MERGE INSTRUCTIONS (for the person or AI combining the report)", bold=True, size=9.5)
for line in [
    "Author of this part: Heinrich Niebuhr (u22555855). Scope: the ordinary-GP and structure-based GP (SBGP) bots built in Java/ECJ, and the SBGP + Adapter-RL hybrid built on them.",
    "Teammates built their own bots, which are described in their own sections. Nothing in this part replaces or evaluates their work.",
    "Every heading, table and figure in this part is numbered with the prefix H so that it can sit beside the other members' parts without clashes.",
    "Yellow [[TODO: ...]] = content still missing. Yellow [[VERIFY: ...]] = a fact to check against code or papers before submission.",
    "Sections H.2.1 and H.6.1 may overlap with shared background or methodology; if so, shorten them to a cross-reference.",
]:
    runs(c.add_paragraph(), line, size=9.5)
doc.add_paragraph()

H("Part H: Genetic Programming, Structure-Based GP and Adapter-RL for microRTS (Heinrich Niebuhr)", 1)

# ======================= H.1
H("H.1 Overview of this part", 2)
P("This part describes one of the bot pipelines developed in the project. We built a genetic programming (GP) system in Java on top of the ECJ library [[VERIFY: ref ECJ v27]] that evolves a complete microRTS agent: one evolved program controls resource gathering, production and combat. Using the same machinery we produced three configurations: an ordinary GP baseline, a structure-based GP variant following Scheepers and Pillay [[ref 18]], and a hybrid in which a trained Adapter-RL module (Jin, Slabaugh and Lucas, 2023) adjusts the decisions of a frozen SBGP bot.")
P("All three configurations share the same representation, primitives, fitness function, training maps and evaluation harness. The SBGP configuration changes only the breeding procedure, and the hybrid only adds a learned correction on top of an unchanged evolved bot. Any difference in results can therefore be attributed to the method being compared rather than to differences in setup.")
P(f"Our main findings are that ordinary GP and SBGP reach the same mean performance but SBGP runs are significantly more consistent, and that adding the adapter to the SBGP bots improves their win rate in {sum(d > 0.05 for d in diffs)} of {len(diffs)} paired runs, by {st.mean(diffs):.1f} percentage points on average.")
FIG(DG["configs"], 6.0, "Figure H.1.", "The three configurations built in this part. Ordinary GP and SBGP share every component except breeding; the hybrid wraps a frozen SBGP bot with a learned correction. All three are evaluated on the same 9 unseen maps against the same 10 opponents.")

# ======================= H.2
H("H.2 Design of the evolved bot", 2)
P("The design was developed over a series of earlier experiments, which we call the development phase. These used a different map split, smaller runs (population 500, 100 generations) and a different evaluation set. Development results quoted in Sections H.2 to H.4 explain design decisions; they are not comparable with the final results in Section H.7.")

H("H.2.1 The decision problem in microRTS", 3)
NOTE("[[Shorten to a cross-reference if the shared background already covers microRTS.]]")
P("In microRTS each player controls a set of units: bases, barracks, workers and three combat types (Light, Heavy and Ranged). Actions are durative, meaning a unit that starts an action is busy for several game cycles. In every cycle the agent must therefore choose an action for each of its idle units. The engine supplies the legal actions for each unit: waiting (NONE), moving, harvesting, returning resources, producing a unit and attacking. The game is deterministic and fully observable, and an agent has a budget of 100 ms per cycle to decide.")

H("H.2.2 Representation: score every action, take the best", 3)
P("The central design decision is the representation. A common way to evolve game agents is to evolve decision rules of the form “if the unit is a worker, then harvest”. We instead evolve a scoring function. For every idle unit, the bot lists the legal actions the engine provides, evaluates the evolved program once per action to give that action a numeric score, and selects the highest-scoring action that the player can afford (Figure H.2).")
FIG(DG["decision"], 5.6, "Figure H.2.", "The score-then-argmax decision loop, run every game cycle. The evolved tree scores each legal action of an idle unit; the highest-scoring affordable action is executed. The per-unit score vector is the point where the adapter later attaches.")
P("This design has three advantages. First, every decision the bot can make is under evolutionary control: the program does not need to know which actions exist, it only has to rank them. Second, the bot can never issue an illegal action, because it only ranks actions the engine already permits. Third, the per-unit score vector is a ranking over actions, which is the form of input the adapter builds on (Section H.4). We chose this representation early, in a hand-written reference bot (Chimera), specifically so that an evolved tree could later replace the hand-written scoring function without changing anything else.")

H("H.2.3 Primitive set", 3)
P("The evolved programs are expression trees. The functions are addition, subtraction, multiplication, protected division (which returns 1 when the divisor is zero), min, max, and a four-argument conditional if>(a, b, c, d) that returns c when a > b and d otherwise. An ephemeral random constant in the range [−10, 10] is also available. Each terminal is a feature of the triple (unit, candidate action, game state) and returns a number. The full set holds 55 entries [[VERIFY: 55 = functions + ERC + terminals in the team params]], grouped in Table H.1.")
TABLE("Table H.1.", "Terminal groups, with examples.", ["Group", "Example terminals", "What they describe"], [
    ["Action type", "aIsMove, aIsAttack, aIsHarvest, aIsNone", "Which kind of action is being scored"],
    ["Production target", "aProducesWorker, aProducesLight, aProducesBarracks", "What a PRODUCE action would build"],
    ["Attack target", "tgtKillable, tgtIsWorker, tgtIsBuilding", "Properties of the unit an attack would hit"],
    ["Movement direction", "movesTowardEnemy, movesTowardEnemyBase, movesTowardResource, movesTowardEnemyMajor", "Whether a MOVE brings the unit closer to a target"],
    ["Placement", "placesTowardEnemy, placesTowardResource", "On which side a new unit would be placed"],
    ["Acting unit", "uIsWorker, uIsBase, uCarrying, uHpFrac", "The unit choosing the action"],
    ["Economy and army", "myResources, myWorkerCount, enemyCombatCount", "The state of both players"],
    ["Spatial and time", "distNearestEnemy, distEnemyBase, threatAtDest, gameTimeK", "Distances, danger and game time"],
])
P("The most important lesson from designing this set concerns which terminals can actually influence a decision. Because the bot takes the argmax over one unit's actions, a term that has the same value for all of that unit's candidate actions cannot change which action wins (Figure H.3). For example, distEnemyBase measures the unit's distance to the enemy base, so it has the same value whichever move the unit considers. It can still be useful as a context gate inside if>, but on its own it cannot steer the unit.")
FIG(DG["terminals"], 6.0, "Figure H.3.", "Why a terminal must vary across a unit's actions to matter. Worker W considers four moves towards enemy E. movesTowardEnemy differs between the moves and can change which one wins; distEnemyBase is the same for every move and cannot.")
P("In our first primitive set every terminal was constant across the four movement directions, so movement was never under evolutionary control: ties between moves were broken by the engine's fixed action order. Adding direction-sensitive terminals such as movesTowardEnemy fixed this immediately: the best individual of the next run beat 13 of the 18 scripted engine bots on its single training map, where previously no individual had beaten the moving opponents reliably. That run's fitness had saturated, so this individual came from the random initial population rather than from selection; the result shows what the new terminals made possible, not what evolution achieved.")

H("H.2.4 One tree per unit role", 3)
P("Each individual holds four trees, one for each unit role: worker, base, barracks and combat unit (Figure H.4). When a unit decides, the bot evaluates the tree for that unit's role. This lets evolution develop different behaviour for each role without first having to learn to distinguish the roles inside one large tree. In practice the barracks and combat trees rarely run, because none of the evolved bots we inspected ever built a barracks (Section H.7.5).")
FIG(DG["roles"], 5.6, "Figure H.4.", "Role trees. Each individual has four trees; the role of the deciding unit selects which tree scores its actions.")

H("H.2.5 Fitness function", 3)
P("The fitness of an individual is computed by playing it against the training opponents on the training maps, on both sides of the map. Lower fitness is better. Each game is scored inside a band determined by its outcome (Figure H.5): a win scores between 0 and 1, with faster wins scoring lower; a draw scores between 1 and 2; and a loss between 2 and 3. Inside the draw and loss bands, the position is set by the quality of play: the normalised value of microRTS's built-in SimpleSqrtEvaluationFunction3, sampled every 50 cycles during the game and averaged. Inside the loss band, surviving longer also lowers the score.")
FIG(DG["bands"], 5.6, "Figure H.5.", "Fitness bands for one game. Each outcome owns a full unit interval, so individuals can be ranked within wins, within draws and within losses.")
P("The fitness of an individual is the mean per-game score multiplied by (1 − b·w), where w is the individual's fraction of games won and b is a win bonus in [0, 1) [[VERIFY: value of win-bonus in the team params]]. RandomBiasedAI, the most random opponent, is played twice per map and side to reduce noise; every other training opponent is played once per map and side, giving 7 opponent slots × 8 maps × 2 sides = 112 games per individual. The evaluation term is normalised to a calibrated range; in the final runs the lowest observed values (about 0.03) fell slightly below the calibrated lower bound (0.05), so values below that bound were clipped; how often this happened was not logged. The effect was the same for both methods.")
P("Each part of this design responds to a problem we observed in earlier versions, where the fitness landscape became flat and selection could not distinguish individuals. A player loses when it has no units left, so every losing final position looks the same, and scoring the final state gave every losing individual the same fitness; sampling the evaluation during play fixed this. In an earlier additive formula the shaping term had little range compared with the gap between outcomes, so different draws were nearly impossible to rank; giving each outcome a full band fixed this. Finally, a subtractive win bonus pushed every all-winning individual below zero, where it was clamped, so fast and slow wins scored the same; making the bonus multiplicative fixed this. The general lesson is that any term that can saturate will saturate exactly where the population ends up, which is also where selection needs a gradient.")

H("H.2.6 Evolutionary parameters", 3)
P("Table H.2 lists the parameters used in the final experiments. Parameters marked as team protocol were agreed by the whole team.")
TABLE("Table H.2.", "GP parameters in the final experiments.", ["Parameter", "Value"], [
    ["Library", "ECJ v27; Java 21 (laptop) and Java 24 (CHPC cluster)"],
    ["Population size", "1200 (team protocol)"],
    ["Generations", "200 (team protocol)"],
    ["Maximum tree depth", "15 (team protocol)"],
    ["Game length cap", "2000 cycles (team protocol)"],
    ["Initialisation", "Ramped half-and-half [[VERIFY]]"],
    ["Selection", "Double tournament for parsimony: fitness tournament of size 7, size tournament of size 1.4"],
    ["Operators", "Subtree crossover, subtree mutation (rate 0.1), reproduction [[VERIFY: crossover and reproduction rates]]"],
    ["Trees per individual", "4 (one per unit role)"],
    ["Training games per individual", "112"],
    ["Duration of the wait (NONE) action", "1 cycle (Section H.2.7)"],
    ["Evaluation threads", "8"],
])
P("Parsimony pressure (the size tournament) was applied identically to ordinary GP and SBGP, so that it cannot explain any difference between them. In development, introducing it coincided with better generalisation: before it, one run produced a 514-node tree that won its own 16×16 training maps but no unseen 16×16 game, while afterwards trees stayed small and began to win unseen 16×16 games. Two training maps were added at the same time, so the improvement cannot be attributed to parsimony alone.")

H("H.2.7 Two properties of the engine that shape the bots", 3)
P("Two properties of the microRTS engine turned out to influence the evolved behaviour strongly. We report them because they affect how the results should be read.")
PB("The engine's action order breaks ties. ", "The engine lists a unit's legal actions in a fixed order: attacks first, then harvest, return, produce, move, and finally wait. When an evolved tree gives several actions the same score, our argmax keeps the first, so the engine's order decides. Part of every evolved bot's behaviour therefore comes from this order rather than from evolution; “attack when possible”, for example, often arises this way. We measured the size of the effect by replacing list order with random tie-breaking for one development bot: it was worth about 3 of 48 games on the validation maps. That bot had been selected on the same maps, so part of this margin may be selection luck.")
PB("The default wait lasts ten cycles. ", "The engine's default wait action keeps a unit idle for ten cycles, while a worker attack takes five. A bot that waits is therefore frozen when the enemy arrives, and since workers have one hit point, whoever strikes first wins worker fights. The scripted engine bots never lock their units this way. Reducing the wait to one cycle was the change that made the difference against WorkerRush on 16×16 maps. In our tests on those maps, no scripted engine bot beat WorkerRush from both sides: the best (WorkerRush's mirror and WorkerRushPlusPlus) won only as player 0, and hand-written scorers over our terminals won at most 3 of 12 games against WorkerRush even with the shorter wait.")

# ======================= H.3
H("H.3 Structure-based genetic programming", 2)
H("H.3.1 The method", 3)
P("Structure-based GP, as proposed by Scheepers and Pillay [[ref 18]] (Genetic Programming and Evolvable Machines 27(1), article 2, 2026), alternates between a global search and a local search over the structure of the programs (Figure H.6). The run starts with an ordinary GP search for Sg generations. The top Dg levels of the best individual found are then stored as a global area. During the following local search, this area is frozen in every individual, so genetic operators can only change the program below level Dg. The local search ends when the best fitness has stopped improving for Wg generations.")
P("A new global search then begins, in which operators act only on the top Dg levels. In this phase, any offspring whose top structure is too similar to a previously explored area is discarded. The aim is to make the search explore new regions of the program space instead of repeatedly refining the same one.")
FIG(DG["sbgp"], 6.0, "Figure H.6.", "The SBGP phase cycle. After a first canonical global search, local and global searches alternate; each global search may only keep offspring that are dissimilar to every earlier global area.")
P("Similarity between two trees is counted top-down from the root: a node counts only if it and all of its ancestors match (Figure H.7). We implemented this index and tested it against the worked example in the paper, where the index is 3.")
FIG(DG["similarity"], 5.8, "Figure H.7.", "The top-down similarity index, illustrated. The root, its left child and that child's left child match together with all their ancestors, giving a similarity of 3; the two mismatched nodes are not counted.")

H("H.3.2 Implementation", 3)
P("SBGP was implemented as two ECJ components: a breeder that controls the phase cycle and filters offspring by similarity, and a node selector that restricts crossover and mutation points to the permitted levels (below the area in local search, inside it in global search). Fitness, primitives, role trees, maps, opponents and parsimony are inherited unchanged from the ordinary-GP configuration, and SBGP bots play through exactly the same bot code. The two methods differ only in breeding. Small unit tests check the similarity index, including the paper's example, and the grafting of an area onto a tree.")

H("H.3.3 Interpretations of the method", 3)
P("Applying the method to our setting required several decisions that the paper does not settle explicitly. Each is a parameter of our implementation; Table H.3 lists them with our reasons. The other SBGP parameters follow the paper: Sg = 10, Tg = 6 and Wg = 10. Because Dg = 2 and our largest function has four arguments, an area holds at most five nodes, so the cap on Tg (Table H.3) always applies. In effect, an offspring is rejected when its whole top area matches an earlier area exactly.")
TABLE("Table H.3.", "Interpretations made when applying SBGP.", ["Decision", "Our choice", "Reason"], [
    ["Depth of the global area Dg", "2 (paper: 4)", "Parsimony keeps our trees at 3–5 levels. With Dg = 4 the frozen area covered whole trees, so local search had almost nothing left to change (about 470 of 500 node selections per generation found no valid point). With Dg = 2, local search produced real gains in pilot runs. Our choice, after discussing it with Prof. Pillay."],
    ["Stagnation", "No improvement greater than 0.001 for Wg generations", "Our fitness is noisy (stochastic opponents), so “no change at all” would rarely occur."],
    ["Entering local search", "Graft the area onto each individual, keeping its own subtrees below the area where they fit", "Preserves the population's diversity below the area."],
    ["Several trees per individual", "One area per role tree; similarity checked only on trees that actually run (worker and base); rejected if any of them is too similar", "In pilot tests, other combinations either never escaped the old area or changed only trees that never run."],
    ["Terminals in the similarity count", "Counted", "Follows the journal text; our trees often have terminals near the root."],
    ["Small areas", "Threshold Tg capped at the size of the area", "Otherwise an area with fewer than Tg nodes can never be matched and the filter silently does nothing."],
    ["Best individual during global search", "Taken only from offspring that pass the filter", "Without this, a pilot run selected the old area again as the “new” one."],
])
PB("Observed behaviour. ", "In pilot runs and in every final run, the first generation of each global search rejected nearly all offspring, which is expected, because the whole population still carries the area just frozen. Rejections fell to about 80% in the next generation and to a few percent after that, so the search did move to new areas. In one pilot run a global search escaped to a better area than the one where local search had stalled, which is the behaviour the method is designed to produce.")

# ======================= H.4
H("H.4 The SBGP + Adapter-RL hybrid", 2)
H("H.4.1 Adapter-RL", 3)
P("Adapter-RL (Jin, Slabaugh and Lucas, 2023) improves an existing agent without changing it. The base agent proposes an action, a neural network called the adapter outputs an adjustment for every possible action, and the two are combined before an action is chosen. For a base action written as a one-hot vector, the combined distribution is")
EQ()
P("where τ is a temperature. A small τ makes the base agent's choice dominant, and a larger τ lets the adapter override it more easily. The adapter is trained with proximal policy optimisation (PPO) while the base agent stays frozen.")

H("H.4.2 Our design", 3)
P("The base agent is a frozen SBGP bot (Section H.3). The design is summarised in Figure H.8, Figure H.9 and Table H.4. The game and the evolved bot run in Java; the adapter runs in Python with PyTorch, connected through JPype so that Python can drive the Java engine directly.")
FIG(DG["dataflow"], 6.3, "Figure H.8.", "Data flow of the hybrid during one decision step. Java encodes the game state and the frozen bot's proposed action; Python combines the base logits with the adapter's adjustments and returns one action code per idle unit, which Java executes. The loop repeats every decision step.")
FIG(DG["network"], 6.3, "Figure H.9.", "The adapter network. A fully convolutional trunk produces 86 action logits for every map cell (each unit reads its own cell) and a single value estimate. The policy head starts at zero, so before training the hybrid plays exactly like the base bot.")
TABLE("Table H.4.", "Adapter design.", ["Component", "Choice"], [
    ["Base distribution", "One-hot of the action the SBGP bot actually takes, scaled by 1/τ (as in the paper)"],
    ["Temperature τ", "0.2"],
    ["Action vocabulary", "86 codes per unit: wait (1), move (4 directions), harvest (4), return (4), produce (6 unit types × 4 directions = 24), attack (7 × 7 grid of relative target positions = 49). Illegal codes are masked out."],
    ["Observation", "29 feature planes over the map: ownership, unit type, hit points, carried resources, resource amounts, busy units and their current actions, walls, which cells must decide, the base agent's proposed action type, and resources and game time"],
    ["Network", "Fully convolutional: three 3×3 convolution layers (64 channels), a global-context layer, a 1×1 policy head (86 logits per cell) and a value head; maps padded to 16×16"],
    ["Initialisation", "Policy head initialised to zero, so the untrained hybrid plays exactly like the base bot"],
    ["Reward", "+1 win, −1 loss, 0 draw, plus 0.5 × the change in normalised state evaluation each step"],
    ["PPO", "γ 0.99, GAE λ 0.95, clip 0.2, value coefficient 1, Adam with learning rate 2.5·10⁻⁴ (as in the paper); 8192 samples per iteration (8 games × 1024 steps); 4 epochs × 8 minibatches; entropy coefficient 0.01; gradient clipping 0.5"],
    ["Training budget", "300 iterations, fixed in advance; final checkpoint used, no checkpoint selection"],
    ["Training data", "Our training maps and training opponents only (Section H.6.2)"],
    ["Evaluation", "Argmax of the combined distribution, so play is deterministic and comparable with the GP bots"],
])

H("H.4.3 Departures from the original paper", 3)
P("We departed from the original method in four ways, each for a practical reason. First, Jin et al. use a convolutional network followed by fully connected layers, which assumes a fixed map size; our network is fully convolutional and works per cell, so one network handles 8×8, 12×12 and 16×16 maps. Second, the paper trains one adapter per map, whereas we train one adapter across all training maps and opponents and evaluate it on unseen maps, so our adapter has to generalise. Third, microRTS requires one action for every idle unit, so we treat these as a joint action and sum the log-probabilities over units. Fourth, our τ = 0.2 lies above the range the authors report as working well (1/1000 to 1/10); in short pilot runs (20 iterations), τ = 0.1 produced almost no exploration, because the base bot's logit of 10 was rarely overridden and the adapter barely changed.")

H("H.4.4 Verification", 3)
P("Before training we checked that the hybrid machinery cannot change behaviour by accident. Sending the base bot's own choices back through the environment reproduced the normal game loop on 64 games, with no mismatch over 36,585 decisions, and random action codes never produced an illegal action. Evaluating bots with an empty adapter through our Python evaluator reproduced the Java harness's results: the macro win rates of three bots differed by at most 0.5 percentage points, from stochastic opponents. Against the four deterministic opponents, the wins, draws and losses of a checked bot matched exactly.")

H("H.4.5 What we learned in development", 3)
NOTE("[[Optional; cut if space is short.]]")
P("Before the final experiments we trained adapters on a single development bot under an earlier protocol. Most checkpoints performed worse on validation maps than the base bot while their training win rates rose, a sign of overfitting to the training maps. To judge whether an improvement was real, we built null controls in which the base bot's decisions were perturbed at random at a matched rate. One of three training seeds produced a checkpoint clearly above the null, and in that case the adapter had learned a single opening decision: on which side of the base the first worker is placed. Against the three deterministic worker bots (WorkerRush, WorkerRushPlusPlus and WorkerDefense), that one decision moved the result between 18 and 42 of 48 games, which shows how sensitive these match-ups are. Because that result depended on one seed and on choosing a checkpoint, the final experiments used a fixed training budget, no checkpoint selection and twenty paired runs.")

# ======================= H.5
H("H.5 Tooling and infrastructure", 2)
P("We built several tools to make the results trustworthy and reproducible (Figure H.10).")
BUL([
    "Saved programs. Each run saves its best-of-run individual together with the full parameter set used to produce it, so any bot can be reloaded, replayed and watched later.",
    "Pruner. The pruner simplifies a finished tree for interpretation. It removes a node only if the simplified tree picks exactly the same action at every decision of a set of traced games, and then verifies the result on separate maps. It is a measurement tool and does not change how a bot plays.",
    "Evaluation harness (TeamEval). The harness plays every bot on every evaluation map against every panel opponent, from both sides, and writes one row per game. It checks at start-up that all opponents load and that the opponents marked deterministic really are deterministic, and it can resume an interrupted evaluation.",
    "Compute. Runs were executed on a laptop (8 threads) and on the CHPC Lengau cluster (programme CSCI0806), with each cluster run as a separate batch job. Adapters were trained on CPU on the cluster, at about 2.5 hours per adapter.",
])
FIG(DG["workflow"], 6.3, "Figure H.10.", "Experimental workflow, from configuration file to statistics. Every evolved bot and every hybrid is evaluated on the same 612 games and recorded game by game.")

# ======================= H.6
H("H.6 Experimental setup", 2)
H("H.6.1 Evaluation protocol (team protocol)", 3)
NOTE("[[Cross-reference the shared methodology if it already describes this.]]")
P("All configurations are evaluated on the nine maps agreed by the team: basesWorkers8x8A, basesWorkersBarracks8x8, basesWorkers8x8Obstacle, basesWorkers12x12, basesWorkers12x12F, SixBasesWorkers12x12, melee16x16Mixed12, TwoBasesBarracks16x16 and EightBasesWorkers16x16 [[VERIFY exact map file names]]. The evaluation panel has ten opponents: RandomBiasedAI, WorkerRush, LightRush, HeavyRush, RangedRush, EconomyLightRush, Mayari, Rojo, Droplet and CoacAI.")
P("Each bot plays every map from both sides, with a cap of 2000 cycles per game. Four opponents (WorkerRush, LightRush, HeavyRush and RangedRush) are deterministic, so each of their games is played once; the other six are stochastic, so each of their games is repeated five times. This gives 612 games per bot. The main measure is the macro win rate: the win rate against each opponent, averaged over the ten opponents so that each opponent counts equally. Draws count as non-wins.")

H("H.6.2 Our training and validation maps", 3)
P("None of the nine evaluation maps were used for training. We trained on eight maps and kept seven further maps for validation (Table H.5). The map basesWorkers8x8 was excluded because it is identical to the evaluation map basesWorkers8x8A.")
TABLE("Table H.5.", "Map split for our configurations.", ["Use", "Maps"], [
    ["Training", "basesWorkers8x8C, basesWorkers8x8E, TwoBasesWorkers8x8, FourBasesWorkers8x8, melee8x8Mixed4, melee16x16Mixed8, basesWorkers16x16B, basesWorkers16x16D [[VERIFY names against config/microrts-team.params]]"],
    ["Validation (development only)", "basesWorkers8x8B, basesWorkers8x8G, ThreeBasesWorkers8x8, melee8x8Mixed6, melee8x8light4, basesWorkers16x16C, basesWorkers16x16E"],
    ["Evaluation (reported)", "the nine team maps (Section H.6.1)"],
    ["Excluded", "basesWorkers8x8 (duplicate of an evaluation map)"],
])
P("The training opponents were WorkerRush, WorkerRushPlusPlus, WorkerDefense, LightRush, RandomBiasedAI (played twice per map and side) and EconomyLightRush. None of the competition bots (Mayari, Rojo, Droplet and CoacAI) were used in training. Four of the ten evaluation opponents (WorkerRush, LightRush, RandomBiasedAI and EconomyLightRush) were also training opponents; the other six were never seen in training. EconomyLightRush is distributed inside the Droplet package, but it is a separate scripted bot. Evaluation-map features such as starting barracks, obstacles and the 12×12 size never occur in training, so the evaluation also tests transfer to new map types.")

H("H.6.3 Runs, seeds and statistics", 3)
P(f"Each GP configuration was run {len(gpm)} times with different random seeds: 11 runs on the laptop and 25 on the cluster. Every run's best-of-run individual was evaluated unpruned; no run was selected or discarded. For the hybrid, an adapter was trained for each of the {len(seeds)} SBGP runs of the first cluster batch (seeds 21242–40242), and each hybrid is compared with its own base bot, which makes it a paired comparison. The number of hybrid runs was fixed before any hybrid result was seen.")
P("We used the Mann-Whitney U test to compare the central performance of ordinary GP and SBGP, the Brown-Forsythe test to compare their spread, and the Wilcoxon signed-rank test for the paired base-versus-hybrid comparison, with the Holm correction across the ten opponents for the per-opponent tests. All tests used α = 0.05. RandomBiasedAI seeds its own random-number generator independently of ECJ, so runs are reproducible only in distribution, not exactly. Seeds, configurations, saved programs and per-game results are archived.")

# ======================= H.7
H("H.7 Results", 2)
H("H.7.1 Ordinary GP versus SBGP", 3)
TABLE("Table H.6.", f"Macro win rate across {len(gpm)} runs per method (%).", ["", "Ordinary GP", "SBGP"], [
    ["Mean", f1d(st.mean(gpm)), f1d(st.mean(sbm))],
    ["Standard deviation", f1d(st.stdev(gpm)), f1d(st.stdev(sbm))],
    ["Median", f1d(st.median(gpm)), f1d(st.median(sbm))],
    ["Minimum", f1d(min(gpm)), f1d(min(sbm))],
    ["Maximum", f1d(max(gpm)), f1d(max(sbm))],
])
FIG(F1, 6.3, "Figure H.11.", "Macro win rate of each run, per method. Each dot is one run; the black line marks the mean and the shaded band ±1 standard deviation.")
P(f"The two methods reach the same central performance (Mann-Whitney U = 641, p ≈ 0.94). SBGP runs are, however, significantly less variable (Brown-Forsythe F(1, 70) = 14.42, p = 0.0003). The variance of ordinary GP is {st.variance(gpm) / st.variance(sbm):.1f} times larger, and SBGP's worst run ({min(sbm):.1f}%) is better than ordinary GP's worst ({min(gpm):.1f}%). Both methods' best runs reach about 67%.")
TABLE("Table H.7.", f"Mean win rate per opponent across {len(gpm)} runs per method (%).", ["Opponent", "Ordinary GP", "SBGP"],
      [[show(o), f1d(st.mean(gp[s][o] for s in gp)), f1d(st.mean(sb[s][o] for s in sb))] for o in OPPS])
P("Both methods beat the army-building scripted bots most of the time, including HeavyRush and RangedRush, which were never seen in training, and RandomBiasedAI nearly always. WorkerRush and the four competition bots are the hardest opponents for both methods.")

H("H.7.2 SBGP versus SBGP + adapter", 3)
up = sum(d > 0.05 for d in diffs); eq = sum(abs(d) <= 0.05 for d in diffs); dn = sum(d < -0.05 for d in diffs)
TABLE("Table H.8.", f"Paired comparison of {len(seeds)} SBGP base bots and their adapter-wrapped versions.", ["", "SBGP (base)", "SBGP + adapter"], [
    ["Mean macro win rate", f"{st.mean(bm):.1f}%", f"{st.mean(hm):.1f}%"],
    ["Mean difference", "", f"{st.mean(diffs):+.1f} percentage points"],
    ["Median difference", "", f"{st.median(diffs):+.1f} percentage points"],
    ["Runs improved / unchanged / worse", "", f"{up} / {eq} / {dn}"],
])
FIG(F2, 4.3, "Figure H.12.", "Macro win rate of each SBGP base bot and its adapter-wrapped version. Green lines mark runs that improved, red lines runs that got worse and grey the unchanged run; the dark line connects the means.")
P(f"The adapter improved its base bot in {up} of {len(diffs)} runs (Wilcoxon signed-rank test, n = 19 after dropping the one tie, W = 9, p ≈ 0.0005). The two runs that got worse lost 2.6 and 0.4 percentage points. In the final hybrids the adapter's choice differed from the base bot's on 13.5–24.0% of decisions, so it changes behaviour substantially rather than making occasional corrections.")
pf = lambda p: (f"{p:.3f}" if p < 0.1 else f"{p:.2f}") + (" *" if p < 0.05 else "")
TABLE("Table H.9.", "Mean win rate per opponent for the paired runs (%), with Wilcoxon signed-rank tests (Holm-corrected; * = significant at 0.05).",
      ["Opponent", "Base", "Hybrid", "Change", "Better / equal / worse", "p (Holm)"],
      [[show(o), f1d(float(dl[o]["base_mean"])), f1d(float(dl[o]["hybrid_mean"])), f"{float(dl[o]['delta_mean']):+.1f}",
        f"{dl[o]['n_up']} / {dl[o]['n_equal']} / {dl[o]['n_down']}", pf(float(dl[o]["p_holm"]))] for o in OPPS])
FIG(F3, 6.0, "Figure H.13.", "Mean win rate per opponent for the SBGP base bots and their adapter-wrapped versions, on the nine unseen evaluation maps.")
P("The mean change is positive against all ten opponents and significant against five. The largest gains were against Droplet (+7.7 points), WorkerRush (+7.5) and Mayari (+4.7); Droplet and Mayari were never played during training.")

H("H.7.3 Decision time", 3)
P("The 100 ms per-cycle budget was respected in almost all cases. For the GP bots, 68 of the 72 evaluated runs never exceeded 90 ms in any cycle. The other four reached 115–231 ms in their single worst cycle; these evaluations ran 24 games in parallel on one cluster node, so this is a worst case under heavy load.")
P("For the adapter we timed three hybrids, single-threaded, against WorkerRush on the nine evaluation maps. The adapter's own work (reading the observation, the forward pass and the argmax) added on average 2.6–3.0 ms per decision step, with a 99th percentile of at most 4.4 ms; this excludes the base bot's scoring and the engine. Of 22,237 timed steps, 7 (0.03%) exceeded 100 ms. The first call of each process takes about 2.6 s while the model loads, which would exceed the budget on the first cycle of a real competition game unless the model is loaded beforehand. The evaluation harness did not enforce the 100 ms budget for any bot.")

H("H.7.4 Training cost", 3)
P("Evolution runs took about 6–8 hours on the cluster (8 cores per run) and between 1.5 and 8.5 hours on the laptop; laptop times are partly inflated because some runs shared the machine. Each adapter took about 2.5 hours of training on 10–12 CPU cores. [[TODO: exact distributions per platform; decide whether to report the cluster only.]]")

H("H.7.5 What the evolved bots do", 3)
NOTE("[[These two examples come from the development phase (an earlier map split, population 500, 100 generations). They illustrate the strategies evolution finds; they are not the bots evaluated above.]]")
P("Because the trees are small and can be pruned without changing play, we can read what the bots do. One SBGP development bot, pruned from 46 to 14 nodes, has the worker tree")
CODE("tgtKillable + movesTowardEnemy + movesTowardEnemyMajor")
P("Every worker charges the nearest enemy and attacks any adjacent target, preferring ones it can kill; once next to the enemy base, the engine's action order makes it attack the base. Its base tree waits one cycle and then, because all actions tie, the engine's action order makes it train workers continuously. This bot never harvests and never builds an army. An ordinary-GP development bot, pruned from 54 to 34 nodes, does the same with a resource-dependent switch.")
P("Watching the bots against WorkerRush showed why they win. WorkerRush keeps one worker at home to harvest, while the evolved bots commit every worker to the attack. With one-hit-point workers, the extra fighter at each clash decides the game. Evolution in effect rediscovered WorkerRush's own strategy, in a more committed and better-timed form.")

# ======================= H.8
H("H.8 Discussion", 2)
PB("Ordinary GP versus SBGP. ", "SBGP did not raise average performance, but it made performance more reliable. This matches the motivation of the method: the global search is meant to escape poor regions of the search space, so it should reduce the number of bad runs more than it raises the best ones. In development we saw ordinary GP fail in two ways: some runs collapse to trivially small trees under parsimony, and others bloat and overfit. SBGP's narrower spread suggests it suffers less from these failures. For a practitioner who can afford only a few runs, a method with a reliable floor is valuable even at an equal mean.")
PB("The adapter. ", "The adapter gave a consistent improvement on top of a frozen evolved bot, on unseen maps, against both training and unseen opponents, with a budget fixed in advance and no checkpoint selection. This is encouraging, because the original paper trained one adapter per map, and RL agents are known to generalise poorly across microRTS maps (Huang and Ontañón, 2021). Several differences from our earlier single-bot experiments, which gave no reliable gain, may explain this: the final training maps were more varied (multi-base and melee maps), twenty different base bots were adapted instead of one, the training budget was fixed in advance, and the outcome measure averaged over ten opponents instead of 48 worker-bot games in which a single opening decision could swing the result between 18 and 42 wins. What the final adapters learned has not yet been analysed; in development, the one successful adapter had learned an opening decision. [[TODO, optional: analyse which action types the final adapters override most.]]")
PB("What the bots rely on. ", "The strongest strategies found are aggressive worker rushes, and part of their behaviour comes from the engine's action order rather than from evolution. This does not bias the comparison between ordinary GP and SBGP, since both are affected equally; the hybrid differs, because its adapter can override decisions the base bot left to list order. It also means the evolved trees are smaller and simpler than the behaviour they produce, and it suggests that evolution exploited the short 2000-cycle games and the small maps, where a rush is strong.")
PB("Comparison with other work. ", "[[TODO: comparison with about three published results (e.g. Huang et al. 2021, Gym-µRTS; Goodfriend 2024; Moraes et al. 2023). Published setups differ in maps, unit table, cycle cap and time budget, so compare as reported results only.]]")
PB("Comparison with the team's other bots. ", "[[TODO: summary numbers and settings from the other members' configurations. State protocol differences neutrally, e.g. which maps each configuration trained on.]]")

# ======================= H.9
H("H.9 Limitations", 2)
BUL([
    "Unit table and rules. We used the engine's original unit table and a 2000-cycle cap. The 2025 microRTS competition used the fine-tuned unit table with a different rule for move conflicts and longer games on 16×16 maps, so our results against the competition bots are not directly comparable with competition results.",
    "Reproducibility. Because RandomBiasedAI seeds itself, runs are reproducible only in distribution.",
    "Different machines. Runs were spread over two platforms with different Java versions and hardware, so training times are reported per platform.",
    f"The hybrid comparison covers {len(seeds)} runs. The hybrids were built on the 20 SBGP runs of the first cluster batch only. Comparing the hybrid's mean with the {len(gpm)}-run GP means is descriptive only, because those runs are not paired with the hybrids.",
    "Evaluation budget. Both methods used the same nominal budget (generations × population size), but the actual number of games played was not logged.",
    "Engine action order. Part of each bot's behaviour is determined by the engine's fixed action order (Section H.2.7).",
])

# ======================= H.10
H("H.10 Conclusion of this part", 2)
P(f"We built a Java GP system that evolves complete microRTS agents by scoring every legal action with an evolved program, and on top of it implemented structure-based GP and an Adapter-RL hybrid. In {len(gpm)} runs per method on nine unseen maps against ten opponents, SBGP matched ordinary GP's mean win rate (about {st.mean(gpm):.0f}%) with significantly less variation between runs. Adding a PPO-trained adapter to frozen SBGP bots raised their mean win rate from {st.mean(bm):.1f}% to {st.mean(hm):.1f}%, improving {up} of {len(diffs)} runs, including against competition bots never seen in training. Along the way we found that the fitness landscape and the primitive set mattered more than any single parameter: a term that cannot differ between a unit's actions cannot steer it, and a fitness term that saturates removes the gradient exactly where it is needed.")

# ======================= H.11
H("H.11 References used in this part", 2)
NOTE("[[Merge into the shared reference list.]]")
BUL([
    "Jin, Slabaugh and Lucas (2023). ADAPTER-RL: Adaptation of Any Agent using Reinforcement Learning. arXiv:2311.11537.",
    "Scheepers and Pillay (2026). A structure-based genetic programming generation constructive hyper-heuristic with transfer learning for combinatorial optimisation. Genetic Programming and Evolvable Machines 27(1), article 2. doi:10.1007/s10710-025-09528-3.",
    "Scheepers (2024). MSc dissertation, University of Pretoria. hdl.handle.net/2263/100614.",
    "Luke, S. ECJ: a Java-based evolutionary computation research system, version 27. [[VERIFY citation form]]",
    "Ontañón, S. (2013). The combinatorial multi-armed bandit problem and its application to real-time strategy games. Proceedings of the AAAI Conference on Artificial Intelligence and Interactive Digital Entertainment (AIIDE 2013). Engine: github.com/santiontanon/microrts.",
    "Schulman et al. (2017). Proximal Policy Optimization Algorithms. arXiv:1707.06347.",
    "Schulman et al. (2015). High-dimensional continuous control using generalized advantage estimation. arXiv:1506.02438.",
    "Huang and Ontañón (2021). Measuring generalization of deep reinforcement learning applied to real-time strategy games. AAAI 2021 Workshop on Reinforcement Learning in Games. [[VERIFY against the Adapter-RL reference list]]",
    "Luke and Panait (2002). Fighting bloat with nonparametric parsimony pressure. Parallel Problem Solving from Nature (PPSN VII).",
    "Koza, J. R. (1992). Genetic Programming: On the Programming of Computers by Means of Natural Selection. MIT Press.",
    "Mann and Whitney (1947). On a test of whether one of two random variables is stochastically larger than the other. Annals of Mathematical Statistics 18(1).",
    "Brown and Forsythe (1974). Robust tests for the equality of variances. Journal of the American Statistical Association 69(346).",
    "Wilcoxon, F. (1945). Individual comparisons by ranking methods. Biometrics Bulletin 1(6).",
    "Holm, S. (1979). A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics 6(2).",
    "JPype: a Python-to-Java bridge. jpype.readthedocs.io. [[VERIFY version used: 1.7.1 on the laptop]]",
])

out = os.path.join(OUT, "part-H-heinrich.docx")
doc.save(out)
print("wrote", out)
print("diagrams in", DIA, ":", ", ".join(os.path.basename(p) for p in DG.values()))